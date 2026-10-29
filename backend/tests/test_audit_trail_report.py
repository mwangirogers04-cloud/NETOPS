from datetime import datetime, timezone
from io import BytesIO
import openpyxl
import pytest
from fastapi.testclient import TestClient

from backend.app.main import app
from backend.app.models.device import Device
from backend.app.models.event import Event
from backend.app.models.user import User, UserSession
from backend.app.services.auth_service import create_access_token
from backend.app.services.event_service import create_event
from backend.app.database import get_db
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from backend.app.database import Base


TEST_DB_URL = "sqlite:///:memory:"


@pytest.fixture
def test_db():
    engine = create_engine(
        TEST_DB_URL,
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(bind=engine)
    TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    db = TestingSessionLocal()
    try:
        yield db
    finally:
        db.close()
        Base.metadata.drop_all(bind=engine)


@pytest.fixture
def auth_headers(test_db):
    admin = User(
        username="superadmin_audit",
        email="superadmin@netops.local",
        hashed_password="fakehashedpassword",
        role="super_admin",
        is_active=True,
        is_verified=True,
    )
    test_db.add(admin)
    test_db.commit()
    test_db.refresh(admin)

    token = create_access_token(data={"sub": admin.username})
    return {"Authorization": f"Bearer {token}"}


def test_audit_trail_comprehensive_attribution(test_db, auth_headers):
    """
    Test that all administrative and security actions (who triggered, deleted,
    added, or changed anything) are pulled into the report with User ID, Role,
    Client IP Address, and timestamp.
    """
    # 1. Create a regular admin user and active user session
    net_admin = User(
        id=42,
        username="netadmin_bob",
        email="bob@netops.local",
        hashed_password="fakehashedpassword",
        role="admin",
        is_active=True,
        is_verified=True,
    )
    test_db.add(net_admin)
    test_db.commit()

    admin_session = UserSession(
        id=10,
        user_id=42,
        username="netadmin_bob",
        role="admin",
        ip_address="192.168.0.55",
        user_agent="Mozilla/5.0 Audit Agent",
        login_at=datetime(2026, 10, 1, 8, 0, 0),
        last_activity_at=datetime(2026, 10, 1, 17, 0, 0),
        is_active=True,
    )
    test_db.add(admin_session)
    test_db.commit()

    # 2. Populate diverse audit actions:
    # A) VLAN Creation (Switch & VLAN)
    create_event(
        db=test_db,
        event_type="vlan_created",
        message="VLAN 50 ('Engineering') created on switch by admin 'netadmin_bob'.",
        actor="netadmin_bob",
        user_id=42,
        role="admin",
        client_ip="192.168.0.55",
    )
    # B) Device Approval (Device Access & Enforcement)
    create_event(
        db=test_db,
        event_type="device_approved",
        message="Device approved by admin 'netadmin_bob': Workstation-12 (192.168.0.99) MAC AA:BB:CC:DD:EE:FF.",
        actor="netadmin_bob",
        user_id=42,
        role="admin",
        client_ip="192.168.0.55",
    )
    # C) Switch Port Configuration (Switch & VLAN)
    create_event(
        db=test_db,
        event_type="switch_port_configured",
        message="Switch port GigabitEthernet0/12 set to VLAN 50 by admin 'netadmin_bob'.",
        actor="netadmin_bob",
        user_id=42,
        role="admin",
        client_ip="192.168.0.55",
    )
    # D) Patch Deployment Triggered (Patch Remediation)
    create_event(
        db=test_db,
        event_type="patch_job_created",
        message="Update deployment job #77 (INSTALL_ALL_UPDATES) queued for workstation-01 by admin 'netadmin_bob'.",
        actor="netadmin_bob",
        user_id=42,
        role="admin",
        client_ip="192.168.0.55",
    )
    # E) User Account Deactivated (User Administration)
    create_event(
        db=test_db,
        event_type="user_deactivated",
        message="User account 'contractor_jim' deactivated by admin 'netadmin_bob'.",
        actor="netadmin_bob",
        user_id=42,
        role="admin",
        client_ip="192.168.0.55",
    )
    # F) VLAN Deletion (Switch & VLAN)
    create_event(
        db=test_db,
        event_type="vlan_deleted",
        message="VLAN 70 removed from switch by admin 'netadmin_bob'.",
        actor="netadmin_bob",
        user_id=42,
        role="admin",
        client_ip="192.168.0.55",
    )
    test_db.commit()

    # Override get_db to return our populated sqlite database
    def override_get_db():
        yield test_db

    app.dependency_overrides[get_db] = override_get_db
    client = TestClient(app)

    try:
        # 1. Test JSON report endpoint
        resp = client.get(
            "/reports/generate?report_type=user_activity&period=monthly&date=2026-10-01",
            headers=auth_headers,
        )
        assert resp.status_code == 200
        data = resp.json()
        assert data["report_type"] == "user_activity"
        assert "audit_events" in data
        assert len(data["audit_events"]) == 6

        # Check attribution fields on each audit item
        for ev in data["audit_events"]:
            assert ev["user_id"] == 42
            assert ev["actor"] == "netadmin_bob"
            assert ev["role"] == "admin"
            assert ev["client_ip"] == "192.168.0.55"
            assert ev["created_at"] is not None
            assert len(ev["details"]) > 0
            assert ev["category"] in (
                "Switch & VLAN Automation",
                "Device Access & Enforcement",
                "Patch & Vulnerability Management",
                "User Administration & Auth",
            )

        # Check summary metrics breakdown
        summary = data["summary"]
        assert summary["audit_actions_in_period"] == 6
        assert summary["users_deactivated_in_period"] == 1
        assert "audit_actions_by_category" in summary
        assert summary["audit_actions_by_category"]["Switch & VLAN Automation"] == 3
        assert summary["audit_actions_by_category"]["Device Access & Enforcement"] == 1
        assert summary["audit_actions_by_category"]["Patch & Vulnerability Management"] == 1
        assert summary["audit_actions_by_category"]["User Administration & Auth"] == 1

        # 2. Test Excel export with dedicated Audit Trail sheet
        excel_resp = client.get(
            "/reports/export/excel?report_type=user_activity&period=monthly&date=2026-10-01",
            headers=auth_headers,
        )
        assert excel_resp.status_code == 200
        assert "spreadsheetml.sheet" in excel_resp.headers["content-type"]

        wb = openpyxl.load_workbook(BytesIO(excel_resp.content))
        assert "Audit Trail (Who Did What)" in wb.sheetnames
        ws_audit = wb["Audit Trail (Who Did What)"]

        # Check audit trail sheet headers
        headers = [ws_audit.cell(1, c).value for c in range(1, 10)]
        assert headers == [
            "Event ID",
            "Timestamp (UTC)",
            "Action Type",
            "Category",
            "User ID",
            "Performed By (Actor)",
            "User Role",
            "Client IP Address",
            "Details / What Was Added, Deleted, or Changed",
        ]

        # Check first data row in Excel
        row2 = [ws_audit.cell(2, c).value for c in range(1, 10)]
        assert row2[4] == 42  # User ID
        assert row2[5] == "netadmin_bob"  # Performed By
        assert row2[6] == "ADMIN"  # Role
        assert row2[7] == "192.168.0.55"  # Client IP

        # 3. Test PDF export contains audit trail
        pdf_resp = client.get(
            "/reports/export/pdf?report_type=user_activity&period=monthly&date=2026-10-01",
            headers=auth_headers,
        )
        assert pdf_resp.status_code == 200
        assert "application/pdf" in pdf_resp.headers["content-type"]
        assert len(pdf_resp.content) > 1000

    finally:
        app.dependency_overrides.pop(get_db, None)


def test_admin_blocked_from_user_activity_and_approvals(test_db):
    """
    Test that standard 'admin' role CANNOT generate or export:
    - User Access & Audit Report (user_activity)
    - Approvals Report (approvals)
    And gets HTTP 403 Forbidden.
    """
    admin_user = User(
        username="standard_admin_role",
        email="standard_admin@netops.local",
        hashed_password="fakehashedpassword",
        role="admin",
        is_active=True,
        is_verified=True,
    )
    test_db.add(admin_user)
    test_db.commit()

    token = create_access_token(data={"sub": admin_user.username})
    admin_headers = {"Authorization": f"Bearer {token}"}

    def override_get_db():
        yield test_db

    app.dependency_overrides[get_db] = override_get_db
    client = TestClient(app)

    try:
        # 1. Attempt generate user_activity -> 403 Forbidden
        r1 = client.get(
            "/reports/generate?report_type=user_activity&period=daily&date=2026-10-01",
            headers=admin_headers,
        )
        assert r1.status_code == 403
        assert "Only Super Administrators" in r1.json()["detail"]

        # 2. Attempt generate approvals -> 403 Forbidden
        r2 = client.get(
            "/reports/generate?report_type=approvals&period=daily&date=2026-10-01",
            headers=admin_headers,
        )
        assert r2.status_code == 403
        assert "Only Super Administrators" in r2.json()["detail"]

        # 3. Attempt export excel user_activity -> 403 Forbidden
        r3 = client.get(
            "/reports/export/excel?report_type=user_activity&period=daily&date=2026-10-01",
            headers=admin_headers,
        )
        assert r3.status_code == 403
        assert "Only Super Administrators" in r3.json()["detail"]

        # 4. Attempt export excel approvals -> 403 Forbidden
        r4 = client.get(
            "/reports/export/excel?report_type=approvals&period=daily&date=2026-10-01",
            headers=admin_headers,
        )
        assert r4.status_code == 403
        assert "Only Super Administrators" in r4.json()["detail"]

        # 5. Attempt export pdf user_activity -> 403 Forbidden
        r5 = client.get(
            "/reports/export/pdf?report_type=user_activity&period=daily&date=2026-10-01",
            headers=admin_headers,
        )
        assert r5.status_code == 403
        assert "Only Super Administrators" in r5.json()["detail"]

        # 6. Attempt export pdf approvals -> 403 Forbidden
        r6 = client.get(
            "/reports/export/pdf?report_type=approvals&period=daily&date=2026-10-01",
            headers=admin_headers,
        )
        assert r6.status_code == 403
        assert "Only Super Administrators" in r6.json()["detail"]

        # 7. Standard Admin CAN access operational reports (e.g. device_inventory, network_activity)
        r7 = client.get(
            "/reports/generate?report_type=device_inventory&period=daily&date=2026-10-01",
            headers=admin_headers,
        )
        assert r7.status_code == 200
        assert r7.json()["report_type"] == "device_inventory"

    finally:
        app.dependency_overrides.pop(get_db, None)

