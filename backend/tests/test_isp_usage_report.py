from datetime import datetime
from io import BytesIO
from unittest.mock import MagicMock
import openpyxl
import pytest
from fastapi.testclient import TestClient

from backend.app.api import reports
from backend.app.main import app
from backend.app.models.device import Device
from backend.app.models.vlan_config import VlanConfig
from backend.app.services.firewall.isp_usage_service import (
    generate_isp_usage_csv,
    get_firewall_isp_bandwidth_usage,
)


client = TestClient(app)


def test_get_firewall_isp_bandwidth_usage_unit():
    mock_db = MagicMock()
    dev1 = Device(
        id=1,
        ip_address="192.168.10.15",
        mac_address="00:1A:2B:3C:4D:5E",
        hostname="PC-HR-01",
        device_type="computer",
        vlan_id=10,
        enforcement_status="allowed",
    )
    dev2 = Device(
        id=2,
        ip_address="192.168.20.25",
        mac_address="00:1A:2B:3C:4D:5F",
        hostname="SRV-DATABASE",
        device_type="server",
        vlan_id=50,
        enforcement_status="allowed",
    )
    dev3 = Device(
        id=3,
        ip_address="192.168.90.99",
        mac_address="00:1A:2B:3C:4D:99",
        hostname="ROGUE-LAPTOP",
        device_type="computer",
        vlan_id=90,
        enforcement_status="quarantined",
    )

    vlan_hr = VlanConfig(vlan_id=10, name="HR_Management", subnet="192.168.10.0/24")
    vlan_srv = VlanConfig(vlan_id=50, name="Servers_Infrastructure", subnet="192.168.50.0/24")

    def mock_scalars(query):
        mock_result = MagicMock()
        query_str = str(query)
        if "devices" in query_str.lower() or "device" in query_str.lower():
            mock_result.all.return_value = [dev1, dev2, dev3]
        elif "vlan" in query_str.lower():
            mock_result.all.return_value = [vlan_hr, vlan_srv]
        else:
            mock_result.all.return_value = []
        return mock_result

    mock_db.scalars.side_effect = mock_scalars

    res = get_firewall_isp_bandwidth_usage(mock_db, period="daily")

    assert res["report_type"] == "isp_internet_usage"
    assert "firewall_info" in res
    assert "wan_interface" in res["firewall_info"]
    assert "isp_provider" in res["firewall_info"]
    assert "bandwidth_summary" in res
    assert res["bandwidth_summary"]["total_consumed_gb"] >= 0
    assert len(res["devices"]) == 3
    assert len(res["departments"]) >= 2

    # Verify quarantined host has minimal/blocked traffic
    quarantined = next(d for d in res["devices"] if d["id"] == 3)
    assert quarantined["enforcement_status"] == "quarantined"
    assert quarantined["firewall_pass"] is False
    assert quarantined["sessions_count"] == 0

    # Verify CSV generator with headers on top row
    csv_out = generate_isp_usage_csv(res)
    csv_lines = [l for l in csv_out.splitlines() if l.strip()]
    # Row 1 must have the column headings at the top
    assert csv_lines[0].startswith("Device Hostname,IP Address,MAC Address")
    # Row 2 must have the content inside each column
    assert "PC-HR-01" in csv_lines[1] or "SRV-DATABASE" in csv_lines[1]
    assert "DEPARTMENT BANDWIDTH BREAKDOWN" in csv_out
    assert "ISP GATEWAY & FIREWALL STATUS" in csv_out


def test_api_generate_isp_usage_report():
    mock_db = MagicMock()
    mock_res = MagicMock()
    mock_res.all.return_value = []
    mock_db.scalars.return_value = mock_res

    def override_get_db():
        yield mock_db

    app.dependency_overrides[reports.get_db] = override_get_db
    try:
        response = client.get("/reports/generate?report_type=isp_internet_usage&period=daily&date=2026-10-24")
    finally:
        app.dependency_overrides.pop(reports.get_db, None)

    assert response.status_code == 200
    data = response.json()
    assert data["report_type"] == "isp_internet_usage"
    assert "firewall_info" in data
    assert "bandwidth_summary" in data


def test_api_isp_usage_dedicated_endpoint():
    mock_db = MagicMock()
    mock_res = MagicMock()
    mock_res.all.return_value = []
    mock_db.scalars.return_value = mock_res

    def override_get_db():
        yield mock_db

    app.dependency_overrides[reports.get_db] = override_get_db
    try:
        response = client.get("/reports/isp-usage?period=weekly")
    finally:
        app.dependency_overrides.pop(reports.get_db, None)

    assert response.status_code == 200
    data = response.json()
    assert data["report_type"] == "isp_internet_usage"
    assert data["period"] == "weekly"


def test_api_isp_usage_csv_download():
    mock_db = MagicMock()
    mock_res = MagicMock()
    mock_res.all.return_value = []
    mock_db.scalars.return_value = mock_res

    def override_get_db():
        yield mock_db

    app.dependency_overrides[reports.get_db] = override_get_db
    try:
        response = client.get("/reports/isp-usage/download/csv?period=daily&date=2026-10-24")
    finally:
        app.dependency_overrides.pop(reports.get_db, None)

    assert response.status_code == 200
    assert "text/csv" in response.headers["content-type"]
    assert "attachment; filename=\"netops_isp_internet_usage" in response.headers["content-disposition"]
    assert response.text.startswith("Device Hostname,IP Address,MAC Address")


def test_api_isp_usage_excel_export():
    dev = Device(
        id=1,
        ip_address="192.168.10.50",
        mac_address="AA:BB:CC:DD:EE:01",
        hostname="HOST-01",
        device_type="computer",
        vlan_id=10,
        enforcement_status="allowed",
    )
    mock_db = MagicMock()
    mock_res = MagicMock()
    mock_res.all.return_value = [dev]
    mock_db.scalars.return_value = mock_res

    def override_get_db():
        yield mock_db

    app.dependency_overrides[reports.get_db] = override_get_db
    try:
        response = client.get("/reports/export/excel?report_type=isp_internet_usage&period=daily&date=2026-10-24")
    finally:
        app.dependency_overrides.pop(reports.get_db, None)

    assert response.status_code == 200
    assert "spreadsheetml.sheet" in response.headers["content-type"]

    wb = openpyxl.load_workbook(BytesIO(response.content))
    expected_sheets = [
        "User Bandwidth Usage",
        "Department Breakdown",
        "ISP & Gateway Overview",
        "Daily Trends",
    ]
    for s in expected_sheets:
        assert s in wb.sheetnames
        ws = wb[s]
        assert ws.cell(1, 1).value is not None
        assert ws.freeze_panes == "A2"
    assert wb.active.title == "User Bandwidth Usage"

    # Verify Sheet 3 ISP & Gateway Overview has column headers on top row
    ws_ov = wb["ISP & Gateway Overview"]
    assert ws_ov.cell(1, 1).value == "Active Firewall"
    assert ws_ov.cell(1, 2).value == "Gateway IP"


def test_api_isp_usage_pdf_export():
    mock_db = MagicMock()
    mock_res = MagicMock()
    mock_res.all.return_value = []
    mock_db.scalars.return_value = mock_res

    def override_get_db():
        yield mock_db

    app.dependency_overrides[reports.get_db] = override_get_db
    try:
        response = client.get("/reports/export/pdf?report_type=isp_internet_usage&period=daily&date=2026-10-24")
    finally:
        app.dependency_overrides.pop(reports.get_db, None)

    assert response.status_code == 200
    assert "application/pdf" in response.headers["content-type"]
    assert len(response.content) > 1000
