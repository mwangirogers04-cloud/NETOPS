from datetime import datetime
from io import BytesIO
from unittest.mock import MagicMock
import openpyxl
import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient

from backend.app.api import reports
from backend.app.api.reports import (
    format_seconds,
    get_period_range,
    parse_report_datetime,
    parse_report_date,
)
from backend.app.main import app
from backend.app.models.device import Device
from backend.app.models.event import Event
from backend.app.models.user import User, UserSession


client = TestClient(app)


def test_parse_report_datetime_valid():
    dt = parse_report_datetime("2026-10-01T14:30")
    assert dt == datetime(2026, 10, 1, 14, 30)

    dt_sec = parse_report_datetime("2026-10-01T14:30:45")
    assert dt_sec == datetime(2026, 10, 1, 14, 30, 45)

    dt_space = parse_report_datetime("2026-10-01 14:30")
    assert dt_space == datetime(2026, 10, 1, 14, 30)


def test_parse_report_datetime_invalid():
    with pytest.raises(ValueError):
        parse_report_datetime("not-a-date")

    with pytest.raises(ValueError):
        parse_report_datetime("")

    with pytest.raises(ValueError):
        parse_report_datetime(None)


def test_get_period_range_standard():
    start, end = get_period_range("daily", date="2026-10-01")
    assert start == datetime(2026, 10, 1, 0, 0)
    assert end == datetime(2026, 10, 2, 0, 0)

    start, end = get_period_range("hourly", date="2026-10-01")
    assert start == datetime(2026, 10, 1, 0, 0)
    assert end == datetime(2026, 10, 1, 1, 0)


def test_get_period_range_custom_datetime():
    start, end = get_period_range(
        "custom",
        start_datetime="2026-10-01T08:00",
        end_datetime="2026-10-01T17:00",
    )
    assert start == datetime(2026, 10, 1, 8, 0)
    assert end == datetime(2026, 10, 1, 17, 0)


def test_get_period_range_custom_dates():
    start, end = get_period_range(
        "custom",
        start_date="2026-10-01",
        end_date="2026-10-05",
    )
    assert start == datetime(2026, 10, 1, 0, 0)
    assert end == datetime(2026, 10, 6, 0, 0)


def test_get_period_range_custom_missing_parameters_raises_400():
    # Calling custom period with no dates/datetimes must raise 400, never 500
    with pytest.raises(HTTPException) as exc_info:
        get_period_range("custom")
    assert exc_info.value.status_code == 400

    # Only start provided
    with pytest.raises(HTTPException) as exc_info:
        get_period_range("custom", start_datetime="2026-10-01T08:00")
    assert exc_info.value.status_code == 400

    with pytest.raises(HTTPException) as exc_info:
        get_period_range("custom", start_date="2026-10-01")
    assert exc_info.value.status_code == 400


def test_get_period_range_custom_inverted_range_raises_400():
    with pytest.raises(HTTPException) as exc_info:
        get_period_range(
            "custom",
            start_datetime="2026-10-05T08:00",
            end_datetime="2026-10-01T08:00",
        )
    assert exc_info.value.status_code == 400

    with pytest.raises(HTTPException) as exc_info:
        get_period_range(
            "custom",
            start_date="2026-10-05",
            end_date="2026-10-01",
        )
    assert exc_info.value.status_code == 400


def test_api_generate_report_device_inventory():
    sample_device = Device(
        id=1,
        ip_address="192.168.1.10",
        mac_address="00:11:22:33:44:55",
        hostname="WORKSTATION-1",
        vendor="Dell",
        device_type="computer",
        vlan_id=10,
        status="online",
        approval_status="approved",
        enforcement_status="allowed",
        classification_status="verified",
        switch_name="SW-CORE-01",
        switch_port="Gi1/0/1",
    )

    db = MagicMock()
    db.scalars.return_value.all.return_value = [sample_device]

    def override_get_db():
        yield db

    app.dependency_overrides[reports.get_db] = override_get_db
    try:
        response = client.get(
            "/reports/generate?report_type=device_inventory&period=daily&date=2026-10-01"
        )
    finally:
        app.dependency_overrides.pop(reports.get_db, None)

    assert response.status_code == 200
    data = response.json()
    assert data["report_type"] == "device_inventory"
    assert data["summary"]["total_devices"] == 1
    assert data["summary"]["online_devices"] == 1
    assert len(data["devices"]) == 1


def test_api_generate_report_custom_period_datetime():
    sample_event = Event(
        id=101,
        event_type="device_detected",
        message="New endpoint detected: 192.168.1.50",
        device_id=1,
        created_at=datetime(2026, 10, 1, 10, 0, 0),
    )

    db = MagicMock()
    db.scalars.return_value.all.return_value = [sample_event]

    def override_get_db():
        yield db

    app.dependency_overrides[reports.get_db] = override_get_db
    try:
        response = client.get(
            "/reports/generate?report_type=network_activity&period=custom&start_datetime=2026-10-01T08:00&end_datetime=2026-10-01T18:00"
        )
    finally:
        app.dependency_overrides.pop(reports.get_db, None)

    assert response.status_code == 200
    data = response.json()
    assert data["period"] == "custom"
    assert data["summary"]["total_events"] == 1
    assert len(data["events"]) == 1


def test_api_generate_report_custom_period_missing_params_returns_400():
    response = client.get(
        "/reports/generate?report_type=network_activity&period=custom"
    )
    assert response.status_code == 400
    assert "Custom period requires" in response.json()["detail"]


def test_api_export_excel_custom_period():
    sample_event = Event(
        id=102,
        event_type="access_enforcement",
        message="Internet access allowed for device 192.168.1.50",
        device_id=1,
        created_at=datetime(2026, 10, 1, 11, 0, 0),
    )

    db = MagicMock()
    db.scalars.return_value.all.return_value = [sample_event]

    def override_get_db():
        yield db

    app.dependency_overrides[reports.get_db] = override_get_db
    try:
        response = client.get(
            "/reports/export/excel?report_type=network_activity&period=custom&start_datetime=2026-10-01T08:00&end_datetime=2026-10-01T18:00"
        )
    finally:
        app.dependency_overrides.pop(reports.get_db, None)

    assert response.status_code == 200
    assert "spreadsheetml.sheet" in response.headers["content-type"]
    assert len(response.content) > 0

    wb = openpyxl.load_workbook(BytesIO(response.content))
    ws = wb.active
    assert ws.title == "Network Activity"
    # Row 1 must have the column headings at the top
    headers = [ws.cell(1, c).value for c in range(1, 6)]
    assert headers == ["Event ID", "Event Type", "Message", "Device ID", "Timestamp"]
    # Row 2 must have the content inside each column
    row2 = [ws.cell(2, c).value for c in range(1, 6)]
    assert row2[0] == 102
    assert row2[1] == "access_enforcement"
    assert "Internet access allowed" in str(row2[2])
    assert str(row2[3]) == "1"
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None


def test_api_export_pdf_with_xml_special_characters():
    # Verify that unescaped characters in message like '<', '>', '&' don't crash ReportLab
    sample_event = Event(
        id=103,
        event_type="access_enforcement",
        message="Device <Unknown & Unregistered> blocked from VLAN 90",
        device_id=1,
        created_at=datetime(2026, 10, 1, 12, 0, 0),
    )

    db = MagicMock()
    db.scalars.return_value.all.return_value = [sample_event]

    def override_get_db():
        yield db

    app.dependency_overrides[reports.get_db] = override_get_db
    try:
        response = client.get(
            "/reports/export/pdf?report_type=network_activity&period=custom&start_datetime=2026-10-01T08:00&end_datetime=2026-10-01T18:00"
        )
    finally:
        app.dependency_overrides.pop(reports.get_db, None)

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")


def test_format_seconds():
    assert format_seconds(None) == "0s"
    assert format_seconds(0) == "0s"
    assert format_seconds(-10) == "0s"
    assert format_seconds(45) == "45s"
    assert format_seconds(125) == "2m 5s"
    assert format_seconds(3665) == "1h 1m 5s"


def test_api_generate_report_user_activity():
    sample_user = User(
        id=1,
        username="netadmin",
        email="netadmin@example.com",
        phone_number="+254700000000",
        role="admin",
        is_active=True,
        is_verified=True,
        created_at=datetime(2026, 10, 1, 10, 0, 0),
    )
    sample_session = UserSession(
        id=10,
        user_id=1,
        username="netadmin",
        role="admin",
        ip_address="192.168.1.15",
        user_agent="Mozilla/5.0",
        login_at=datetime(2026, 10, 1, 10, 5, 0),
        last_activity_at=datetime(2026, 10, 1, 11, 10, 0),
        logout_at=datetime(2026, 10, 1, 11, 10, 0),
        duration_seconds=3900,
        is_active=False,
    )

    db = MagicMock()
    mock_res1 = MagicMock()
    mock_res1.all.return_value = [sample_user]
    mock_res2 = MagicMock()
    mock_res2.all.return_value = [sample_user]
    mock_res3 = MagicMock()
    mock_res3.all.return_value = [sample_session]
    db.scalars.side_effect = [mock_res1, mock_res2, mock_res3]

    def override_get_db():
        yield db

    app.dependency_overrides[reports.get_db] = override_get_db
    try:
        response = client.get(
            "/reports/generate?report_type=user_activity&period=daily&date=2026-10-01"
        )
    finally:
        app.dependency_overrides.pop(reports.get_db, None)

    assert response.status_code == 200
    data = response.json()
    assert data["report_type"] == "user_activity"
    assert data["summary"]["total_users"] == 1
    assert data["summary"]["users_created_in_period"] == 1
    assert data["summary"]["total_sessions"] == 1
    assert data["summary"]["total_duration_seconds"] == 3900
    assert data["summary"]["total_duration_formatted"] == "1h 5m 0s"
    assert len(data["users_created"]) == 1
    assert data["users_created"][0]["username"] == "netadmin"
    assert len(data["sessions"]) == 1
    assert data["sessions"][0]["username"] == "netadmin"
    assert data["sessions"][0]["duration_formatted"] == "1h 5m 0s"


def test_api_export_excel_user_activity():
    sample_user = User(
        id=1,
        username="developer",
        email="dev@netops.local",
        phone_number="+254711111111",
        role="super_admin",
        is_active=True,
        is_verified=True,
        created_at=datetime(2026, 10, 1, 8, 0, 0),
    )
    sample_session = UserSession(
        id=1,
        user_id=1,
        username="developer",
        role="super_admin",
        ip_address="192.168.1.10",
        user_agent="Mozilla/5.0",
        login_at=datetime(2026, 10, 1, 8, 30, 0),
        last_activity_at=datetime(2026, 10, 1, 9, 30, 0),
        logout_at=datetime(2026, 10, 1, 9, 30, 0),
        duration_seconds=3600,
        is_active=False,
    )

    db = MagicMock()
    mock_res1 = MagicMock()
    mock_res1.all.return_value = [sample_user]
    mock_res2 = MagicMock()
    mock_res2.all.return_value = [sample_user]
    mock_res3 = MagicMock()
    mock_res3.all.return_value = [sample_session]
    db.scalars.side_effect = [mock_res1, mock_res2, mock_res3]

    def override_get_db():
        yield db

    app.dependency_overrides[reports.get_db] = override_get_db
    try:
        response = client.get(
            "/reports/export/excel?report_type=user_activity&period=daily&date=2026-10-01"
        )
    finally:
        app.dependency_overrides.pop(reports.get_db, None)

    assert response.status_code == 200
    assert "spreadsheetml.sheet" in response.headers["content-type"]
    assert len(response.content) > 0

    wb = openpyxl.load_workbook(BytesIO(response.content))
    assert wb.sheetnames == ["User Sessions", "Users Created", "Report Summary"]

    # 1. User Sessions sheet
    ws_sess = wb["User Sessions"]
    assert wb.active == ws_sess  # Must be active first sheet
    sess_headers = [ws_sess.cell(1, c).value for c in range(1, 11)]
    assert sess_headers == [
        "Session ID",
        "Username",
        "Role",
        "Login Time",
        "Last Activity",
        "Logout Time",
        "Duration",
        "Duration (Seconds)",
        "IP Address",
        "Session Status",
    ]
    sess_row2 = [ws_sess.cell(2, c).value for c in range(1, 11)]
    assert sess_row2[0] == 1
    assert sess_row2[1] == "developer"
    assert sess_row2[2] == "SUPER_ADMIN"
    assert sess_row2[6] == "1h 0m 0s"
    assert sess_row2[7] == 3600
    assert sess_row2[8] == "192.168.1.10"
    assert sess_row2[9] == "Ended"
    assert ws_sess.freeze_panes == "A2"
    assert ws_sess.auto_filter.ref is not None

    # 2. Users Created sheet
    ws_users = wb["Users Created"]
    user_headers = [ws_users.cell(1, c).value for c in range(1, 9)]
    assert user_headers == [
        "User ID",
        "Username",
        "Email",
        "Phone Number",
        "Role",
        "Created At",
        "Verification Status",
        "Account Status",
    ]
    user_row2 = [ws_users.cell(2, c).value for c in range(1, 9)]
    assert user_row2[0] == 1
    assert user_row2[1] == "developer"
    assert user_row2[2] == "dev@netops.local"
    assert user_row2[4] == "SUPER_ADMIN"
    assert user_row2[6] == "Verified"
    assert user_row2[7] == "Active"
    assert ws_users.freeze_panes == "A2"
    assert ws_users.auto_filter.ref is not None

    # 3. Report Summary sheet
    ws_sum = wb["Report Summary"]
    sum_headers = [ws_sum.cell(1, c).value for c in range(1, 3)]
    assert sum_headers == ["Parameter / Metric", "Value"]
    assert ws_sum.cell(2, 1).value == "Report Type"
    assert ws_sum.cell(2, 2).value == "User Access & Audit Report"


def test_api_export_pdf_user_activity():
    sample_user = User(
        id=1,
        username="developer",
        email="dev@netops.local",
        phone_number="+254711111111",
        role="super_admin",
        is_active=True,
        is_verified=True,
        created_at=datetime(2026, 10, 1, 8, 0, 0),
    )
    sample_session = UserSession(
        id=1,
        user_id=1,
        username="developer",
        role="super_admin",
        ip_address="192.168.1.10",
        user_agent="Mozilla/5.0",
        login_at=datetime(2026, 10, 1, 8, 30, 0),
        last_activity_at=datetime(2026, 10, 1, 9, 30, 0),
        logout_at=datetime(2026, 10, 1, 9, 30, 0),
        duration_seconds=3600,
        is_active=False,
    )

    db = MagicMock()
    mock_res1 = MagicMock()
    mock_res1.all.return_value = [sample_user]
    mock_res2 = MagicMock()
    mock_res2.all.return_value = [sample_user]
    mock_res3 = MagicMock()
    mock_res3.all.return_value = [sample_session]
    db.scalars.side_effect = [mock_res1, mock_res2, mock_res3]

    def override_get_db():
        yield db

    app.dependency_overrides[reports.get_db] = override_get_db
    try:
        response = client.get(
            "/reports/export/pdf?report_type=user_activity&period=daily&date=2026-10-01"
        )
    finally:
        app.dependency_overrides.pop(reports.get_db, None)

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")


def test_api_export_excel_device_inventory():
    sample_device = Device(
        id=5,
        ip_address="192.168.1.55",
        mac_address="AA:BB:CC:DD:EE:FF",
        hostname="WORKSTATION-01",
        vendor="Dell Inc.",
        device_type="computer",
        vlan_id=20,
        status="online",
        approval_status="approved",
        enforcement_status="allowed",
        classification_status="classified",
        switch_name="SW-CORE-01",
        switch_port="Gi1/0/12",
    )

    db = MagicMock()
    mock_res = MagicMock()
    mock_res.all.return_value = [sample_device]
    db.scalars.return_value = mock_res

    def override_get_db():
        yield db

    app.dependency_overrides[reports.get_db] = override_get_db
    try:
        response = client.get(
            "/reports/export/excel?report_type=device_inventory&period=daily&date=2026-10-01"
        )
    finally:
        app.dependency_overrides.pop(reports.get_db, None)

    assert response.status_code == 200
    assert "spreadsheetml.sheet" in response.headers["content-type"]

    wb = openpyxl.load_workbook(BytesIO(response.content))
    assert wb.sheetnames == ["Device Inventory", "Report Summary"]
    ws_dev = wb["Device Inventory"]
    assert wb.active == ws_dev

    # Row 1 headers must appear at the top
    headers = [ws_dev.cell(1, c).value for c in range(1, 14)]
    assert headers == [
        "Device ID",
        "IP Address",
        "MAC Address",
        "Hostname",
        "Vendor",
        "Device Type",
        "VLAN ID",
        "Status",
        "Approval Status",
        "Enforcement Status",
        "Classification Status",
        "Switch Name",
        "Switch Port",
    ]

    # Row 2 must have the content inside each column
    row2 = [ws_dev.cell(2, c).value for c in range(1, 14)]
    assert row2[0] == 5
    assert row2[1] == "192.168.1.55"
    assert row2[2] == "AA:BB:CC:DD:EE:FF"
    assert row2[3] == "WORKSTATION-01"
    assert row2[5] == "computer"
    assert row2[6] == 20
    assert ws_dev.freeze_panes == "A2"
    assert ws_dev.auto_filter.ref is not None

