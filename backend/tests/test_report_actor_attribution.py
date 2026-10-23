from datetime import datetime
from io import BytesIO
from unittest.mock import MagicMock
import openpyxl
import pytest
from fastapi.testclient import TestClient

from backend.app.api import reports
from backend.app.main import app
from backend.app.models.device import Device
from backend.app.models.event import Event
from backend.app.models.user import User, UserSession
from backend.app.services.auth_service import create_access_token


client = TestClient(app)


def test_approvals_report_includes_actor_and_deletion():
    # Simulate events with specific actors for device approval, rejection, and deletion
    event_approved = Event(
        id=201,
        event_type="device_approved",
        message="Device approved by admin 'alice': Host1 (192.168.1.10) MAC 00:11:22:33:44:55",
        device_id=1,
        actor="alice",
        created_at=datetime(2026, 10, 1, 10, 0, 0),
    )
    event_rejected = Event(
        id=202,
        event_type="device_rejected",
        message="Device rejected by super_admin 'developer': Host2 (192.168.1.20) MAC 00:11:22:33:44:66",
        device_id=2,
        actor="developer",
        created_at=datetime(2026, 10, 1, 11, 0, 0),
    )
    event_deleted = Event(
        id=203,
        event_type="device_permanently_deleted",
        message="Machine permanently deleted from inventory by admin 'bob': Host3 (192.168.1.30 / 00:11:22:33:44:77)",
        device_id=None,
        actor="bob",
        created_at=datetime(2026, 10, 1, 12, 0, 0),
    )

    db = MagicMock()
    mock_events = MagicMock()
    mock_events.all.return_value = [event_approved, event_rejected, event_deleted]
    db.scalars.return_value = mock_events

    mock_counts = MagicMock()
    mock_counts.all.return_value = [("device_approved", 1), ("device_rejected", 1), ("device_permanently_deleted", 1)]
    db.execute.return_value = mock_counts

    def override_get_db():
        yield db

    app.dependency_overrides[reports.get_db] = override_get_db
    try:
        response = client.get(
            "/reports/generate?report_type=approvals&period=daily&date=2026-10-01"
        )
    finally:
        app.dependency_overrides.pop(reports.get_db, None)

    assert response.status_code == 200
    data = response.json()
    assert data["report_type"] == "approvals"
    assert data["summary"]["total_events"] == 3
    events = data["events"]
    assert len(events) == 3

    # Check that each event knows who did that
    assert events[0]["actor"] == "alice"
    assert "admin 'alice'" in events[0]["message"]
    assert events[1]["actor"] == "developer"
    assert "super_admin 'developer'" in events[1]["message"]
    assert events[2]["actor"] == "bob"
    assert "admin 'bob'" in events[2]["message"]


def test_user_activity_report_includes_created_by_and_audit_trail():
    sample_user = User(
        id=5,
        username="john_doe",
        email="john@example.com",
        phone_number="+254700000001",
        role="user",
        is_active=False,
        is_verified=True,
        created_by="developer",
        created_at=datetime(2026, 10, 1, 9, 0, 0),
    )
    user_created_ev = Event(
        id=301,
        event_type="user_created",
        message="New user 'john_doe' (role: user) created by super_admin 'developer'.",
        actor="developer",
        created_at=datetime(2026, 10, 1, 9, 0, 0),
    )
    user_deactivated_ev = Event(
        id=302,
        event_type="user_deactivated",
        message="User account 'john_doe' deactivated by admin 'netadmin'.",
        actor="netadmin",
        created_at=datetime(2026, 10, 1, 14, 0, 0),
    )
    password_reset_ev = Event(
        id=303,
        event_type="password_reset",
        message="Password for user 'john_doe' directly reset by super_admin 'developer'.",
        actor="developer",
        created_at=datetime(2026, 10, 1, 15, 0, 0),
    )

    db = MagicMock()
    mock_users_created = MagicMock()
    mock_users_created.all.return_value = [sample_user]
    mock_all_users = MagicMock()
    mock_all_users.all.return_value = [sample_user]
    mock_sessions = MagicMock()
    mock_sessions.all.return_value = []
    mock_audit = MagicMock()
    mock_audit.all.return_value = [user_created_ev, user_deactivated_ev, password_reset_ev]

    db.scalars.side_effect = [mock_users_created, mock_all_users, mock_sessions, mock_audit]

    def override_get_db():
        yield db

    app.dependency_overrides[reports.get_db] = override_get_db
    try:
        response = client.get(
            "/reports/generate?report_type=user_activity&period=daily&date=2026-10-01"
        )
    finally:
        app.dependency_overrides.pop(reports.get_db, None)

    assert response.status_code == 200
    data = response.json()
    assert data["report_type"] == "user_activity"

    # Verify created_by attribution
    assert len(data["users_created"]) == 1
    assert data["users_created"][0]["username"] == "john_doe"
    assert data["users_created"][0]["created_by"] == "developer"

    # Verify audit summary
    assert data["summary"]["users_deactivated_in_period"] == 1
    assert data["summary"]["password_resets_in_period"] == 1
    assert data["summary"]["audit_actions_in_period"] == 3

    # Verify audit trail events
    audit_evs = data["audit_events"]
    assert len(audit_evs) == 3
    assert audit_evs[0]["event_type"] == "user_created"
    assert audit_evs[0]["actor"] == "developer"
    assert audit_evs[1]["event_type"] == "user_deactivated"
    assert audit_evs[1]["actor"] == "netadmin"
    assert audit_evs[2]["event_type"] == "password_reset"
    assert audit_evs[2]["actor"] == "developer"


def test_export_excel_with_actors():
    event_approved = Event(
        id=201,
        event_type="device_approved",
        message="Device approved by admin 'alice': Host1",
        device_id=1,
        actor="alice",
        created_at=datetime(2026, 10, 1, 10, 0, 0),
    )
    db = MagicMock()
    mock_events = MagicMock()
    mock_events.all.return_value = [event_approved]
    db.scalars.return_value = mock_events
    mock_counts = MagicMock()
    mock_counts.all.return_value = [("device_approved", 1)]
    db.execute.return_value = mock_counts

    def override_get_db():
        yield db

    app.dependency_overrides[reports.get_db] = override_get_db
    try:
        response = client.get(
            "/reports/export/excel?report_type=approvals&period=daily&date=2026-10-01"
        )
    finally:
        app.dependency_overrides.pop(reports.get_db, None)

    assert response.status_code == 200
    wb = openpyxl.load_workbook(BytesIO(response.content))
    ws = wb.active
    # First 6 headers
    headers = [ws.cell(1, c).value for c in range(1, 7)]
    assert headers == ["Event ID", "Event Type", "Message", "Device ID", "Timestamp", "Performed By"]
    # Row 2 Performed By column
    assert ws.cell(2, 6).value == "alice"


from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from backend.app.database import Base, get_db
from backend.app.services.auth_service import hash_password

TEST_DB_URL = "sqlite:///:memory:"
test_engine = create_engine(
    TEST_DB_URL,
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestingSession = sessionmaker(autocommit=False, autoflush=False, bind=test_engine)


def test_device_actions_record_actor_in_event():
    Base.metadata.create_all(bind=test_engine)
    db = TestingSession()

    def test_get_db():
        yield db

    app.dependency_overrides[get_db] = test_get_db
    try:
        # Create admin user
        admin = User(
            username="admin_sarah",
            email="sarah@netops.local",
            hashed_password=hash_password("Secret123!"),
            role="admin",
            is_active=True,
            is_verified=True,
        )
        device = Device(
            ip_address="192.168.1.100",
            mac_address="00:11:22:33:44:55",
            hostname="PC-SARAH",
            approval_status="pending",
            enforcement_status="quarantined",
            vlan_id=90,
        )
        db.add(admin)
        db.add(device)
        db.commit()
        db.refresh(admin)
        db.refresh(device)

        token = create_access_token({"sub": admin.username})
        headers = {"Authorization": f"Bearer {token}"}

        # 1. Approve device
        res = client.patch(f"/devices/{device.id}/approve", headers=headers)
        assert res.status_code == 200

        approval_ev = db.scalar(
            select(Event).where(Event.event_type == "device_approved").order_by(Event.id.desc())
        )
        assert approval_ev is not None
        assert approval_ev.actor == "admin_sarah"
        assert "admin 'admin_sarah'" in approval_ev.message

        # 2. Reject device
        res = client.patch(f"/devices/{device.id}/reject", headers=headers)
        assert res.status_code == 200

        reject_ev = db.scalar(
            select(Event).where(Event.event_type == "device_rejected").order_by(Event.id.desc())
        )
        assert reject_ev is not None
        assert reject_ev.actor == "admin_sarah"
        assert "admin 'admin_sarah'" in reject_ev.message

        # 3. Permanently delete device
        res = client.delete(f"/devices/{device.id}", headers=headers)
        assert res.status_code == 200

        del_ev = db.scalar(
            select(Event).where(Event.event_type == "device_permanently_deleted").order_by(Event.id.desc())
        )
        assert del_ev is not None
        assert del_ev.actor == "admin_sarah"
        assert "admin 'admin_sarah'" in del_ev.message

    finally:
        app.dependency_overrides.pop(get_db, None)
        db.close()
        Base.metadata.drop_all(bind=test_engine)


def test_user_lifecycle_records_actor_in_event():
    Base.metadata.create_all(bind=test_engine)
    db = TestingSession()

    def test_get_db():
        yield db

    app.dependency_overrides[get_db] = test_get_db
    try:
        # Create super admin
        superadmin = User(
            username="super_tom",
            email="tom@example.com",
            hashed_password=hash_password("SuperSecret123!"),
            role="super_admin",
            is_active=True,
            is_verified=True,
        )
        db.add(superadmin)
        db.commit()

        token = create_access_token({"sub": superadmin.username})
        headers = {"Authorization": f"Bearer {token}"}

        # 1. Create a new user
        create_res = client.post(
            "/users/",
            json={
                "username": "new_engineer",
                "email": "engineer@example.com",
                "phone_number": "+254711223344",
                "password": "Password123!",
                "role": "user",
            },
            headers=headers,
        )
        assert create_res.status_code == 201
        new_user_data = create_res.json()["user"]
        assert new_user_data["created_by"] == "super_tom"

        create_ev = db.scalar(
            select(Event).where(Event.event_type == "user_created").order_by(Event.id.desc())
        )
        assert create_ev is not None
        assert create_ev.actor == "super_tom"
        assert "created by super_admin 'super_tom'" in create_ev.message

        user_id = new_user_data["id"]

        # 2. Deactivate the user
        deact_res = client.patch(
            f"/users/{user_id}",
            json={"is_active": False},
            headers=headers,
        )
        assert deact_res.status_code == 200

        deact_ev = db.scalar(
            select(Event).where(Event.event_type == "user_deactivated").order_by(Event.id.desc())
        )
        assert deact_ev is not None
        assert deact_ev.actor == "super_tom"
        assert "deactivated by super_admin 'super_tom'" in deact_ev.message

        # 3. Direct password reset
        reset_res = client.post(
            "/auth/forgot-password/admin-reset",
            json={
                "admin_identifier": "super_tom",
                "admin_password": "SuperSecret123!",
                "target_identifier": "new_engineer",
                "new_password": "NewSecret456!",
            },
        )
        assert reset_res.status_code == 200

        reset_ev = db.scalar(
            select(Event).where(Event.event_type == "password_reset").order_by(Event.id.desc())
        )
        assert reset_ev is not None
        assert reset_ev.actor == "super_tom"
        assert "directly reset by Developer (Super Admin) 'super_tom'" in reset_ev.message

    finally:
        app.dependency_overrides.pop(get_db, None)
        db.close()
        Base.metadata.drop_all(bind=test_engine)

